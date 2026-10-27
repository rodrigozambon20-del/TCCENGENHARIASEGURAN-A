# -*- coding: utf-8 -*-
"""
Gera Planilha_Dados_Poda.xlsx para lançar os dados do estudo de caso de poda.
Calcula NEN (NHO-01), classificação e resumo de conformidade.

Uso:  python3 gerar_planilha_campo.py      Requer: pip install openpyxl
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

OUT = "Planilha_Dados_Poda.xlsx"
HEAD = PatternFill("solid", fgColor="16233F")
INPUT = PatternFill("solid", fgColor="FFF8DC")
CALC = PatternFill("solid", fgColor="E8EEF6")
WHITE = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="B8C2D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def header(ws, cols, widths, row=1):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = HEAD; cell.font = WHITE; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def paint(ws, r1, r2, c1, c2, fill):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill; cell.border = BOX; cell.alignment = WRAP


# ------------------------------------------------------------ Leia-me -------
ws = wb.active; ws.title = "Leia-me"
ws.column_dimensions["A"].width = 110
lines = [
    "PLANILHA DE DADOS — ESTUDO DE CASO: PODA JUNTO À REDE DE DISTRIBUIÇÃO",
    "",
    "Células AMARELAS = você preenche.  Células AZUIS = calculadas (não editar).",
    "Nunca escreva nomes: use P1, P2 para os podadores e o nº de série dos equipamentos.",
    "",
    "Abas:",
    "  Dosimetria   — uma linha por jornada medida (ficha F2). Calcula NEN e classificação.",
    "  Atividades   — diário da jornada (ficha F3), para explicar a dose.",
    "  Verificacao  — itens das fichas F4 (NR-10), F5 (NR-35) e F6 (NR-12): C / NC / NA.",
    "  EPI          — ficha F7.",
    "  Fotos        — ficha F9.",
    "  Resumo       — totais automáticos para o capítulo de resultados.",
    "  Referencia   — parâmetros e critérios usados nos cálculos.",
    "",
    "Fórmula do NEN (NHO-01, q = 3): NEN = 85 + 10·log10(D/100), com D = dose da jornada (%).",
    "Se a medição cobriu só parte da jornada: D projetada = D medida × (jornada / TE).",
    "Amostra válida: diferença entre calibração antes e depois ≤ ±1 dB.",
]
for i, t in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t)
ws["A1"].font = Font(bold=True, size=13, color="16233F")

# ----------------------------------------------------------- Dosimetria ------
ws = wb.create_sheet("Dosimetria")
cols = ["Amostra", "Data", "Podador", "Atividade predominante", "Dosímetro (nº série)",
        "Calib. antes (dB)", "Calib. depois (dB)", "Diferença (dB)", "Amostra válida?",
        "Jornada (min)", "TE medido (min)", "Dose NHO-01 medida (%)", "Dose NR-15 (%)",
        "Leq/Lavg (dB(A))", "Dose NHO-01 projetada (%)", "NEN (dB(A))",
        "Classificação NHO-01", "Situação NR-15 (Anexo 1)", "Protetor auditivo no corte",
        "Observações"]
header(ws, cols, [9, 11, 9, 24, 15, 10, 10, 10, 11, 10, 10, 12, 11, 11, 13, 10, 20, 22, 16, 30])
N = 12
for r in range(2, N + 2):
    ws.cell(row=r, column=1, value=f"A{r-1:02d}")
    ws.cell(row=r, column=8, value=f'=IF(OR(F{r}="",G{r}=""),"",ROUND(G{r}-F{r},1))')
    ws.cell(row=r, column=9, value=f'=IF(H{r}="","",IF(ABS(H{r})<=1,"Sim","NÃO — repetir"))')
    ws.cell(row=r, column=15, value=f'=IF(OR(L{r}="",K{r}="",J{r}=""),"",ROUND(L{r}*J{r}/K{r},1))')
    ws.cell(row=r, column=16, value=f'=IF(OR(O{r}="",O{r}<=0),"",ROUND(85+10*LOG10(O{r}/100),1))')
    ws.cell(row=r, column=17, value=(f'=IF(P{r}="","",IF(I{r}="NÃO — repetir","Amostra inválida",'
                                     f'IF(P{r}>=85,"Acima do limite",'
                                     f'IF(P{r}>=82,"Nível de ação","Aceitável"))))'))
    ws.cell(row=r, column=18, value=(f'=IF(M{r}="","",IF(I{r}="NÃO — repetir","Amostra inválida",'
                                     f'IF(M{r}>100,"Acima do LT (insalubre)","Abaixo do LT")))'))
    ws.cell(row=r, column=2).number_format = "DD/MM/YYYY"
paint(ws, 2, N + 1, 1, 7, INPUT); paint(ws, 2, N + 1, 10, 14, INPUT)
paint(ws, 2, N + 1, 19, 20, INPUT)
paint(ws, 2, N + 1, 8, 9, CALC); paint(ws, 2, N + 1, 15, 18, CALC)
dv = DataValidation(type="list", formula1='"P1,P2,P3"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"C2:C{N+1}")
dvp = DataValidation(type="list", formula1='"Sempre,Às vezes,Não"', allow_blank=True)
ws.add_data_validation(dvp); dvp.add(f"S2:S{N+1}")
red = PatternFill("solid", fgColor="F4C7C3"); amber = PatternFill("solid", fgColor="FCE8B2")
ws.conditional_formatting.add(f"Q2:Q{N+1}", FormulaRule(formula=['Q2="Acima do limite"'], fill=red))
ws.conditional_formatting.add(f"Q2:Q{N+1}", FormulaRule(formula=['Q2="Nível de ação"'], fill=amber))
ws.conditional_formatting.add(f"R2:R{N+1}", FormulaRule(formula=['R2="Acima do LT (insalubre)"'], fill=red))
ws.conditional_formatting.add(f"I2:I{N+1}", FormulaRule(formula=['I2="NÃO — repetir"'], fill=red))

# ----------------------------------------------------------- Atividades ------
ws = wb.create_sheet("Atividades")
cols = ["Amostra", "Início (hh:mm)", "Fim (hh:mm)", "Duração (min)", "Atividade",
        "Equipamento ligado", "Obs."]
header(ws, cols, [9, 12, 12, 12, 36, 20, 30])
for r in range(2, 202):
    ws.cell(row=r, column=4, value=f'=IF(OR(B{r}="",C{r}=""),"",ROUND((C{r}-B{r})*1440,0))')
    ws.cell(row=r, column=2).number_format = "hh:mm"
    ws.cell(row=r, column=3).number_format = "hh:mm"
paint(ws, 2, 201, 1, 3, INPUT); paint(ws, 2, 201, 5, 7, INPUT); paint(ws, 2, 201, 4, 4, CALC)
dva = DataValidation(type="list", formula1='"Motosserra,Motopoda,Podão manual,Nenhum,Veículo"',
                     allow_blank=True)
ws.add_data_validation(dva); dva.add("F2:F201")

# ----------------------------------------------------------- Verificação -----
ws = wb.create_sheet("Verificacao")
cols = ["Ficha", "Norma", "#", "Item verificado", "Data", "Podador/equip.",
        "Resultado (C/NC/NA)", "Nº da foto (evidência)", "Observação"]
header(ws, cols, [7, 16, 5, 60, 11, 14, 12, 14, 34])
F4 = ["Treinamento NR-10 básico (40 h) — certificado",
      "Curso complementar SEP (40 h)",
      "Autorização formal para serviços em/próximo de instalações elétricas",
      "Ordem de serviço / programação do serviço",
      "Análise de risco (APR) antes de iniciar, específica do local",
      "Rede desligada: sequência de desenergização seguida na ordem",
      "Rede energizada: pessoas/ferramentas fora da zona de risco; zona controlada só autorizados",
      "Ferramentas isolantes com ensaio dielétrico em dia",
      "Controle da queda dos galhos (não encostaram na rede)",
      "Comunicação com centro de operação",
      "Serviço não realizado por trabalhador sozinho",
      "Pessoa capacitada em primeiros socorros/RCP e plano de emergência",
      "Área sinalizada e isolada"]
F5 = ["Treinamento NR-35 (mín. 8 h) — certificado",
      "ASO com aptidão para trabalho em altura",
      "Análise de risco do trabalho em altura",
      "Permissão de trabalho (se não rotineiro)",
      "Cinturão e talabarte inspecionados antes do uso",
      "Ancoragem escolhida e verificada",
      "Trabalhador conectado o tempo todo em altura",
      "Meio de acesso adequado (escada estabilizada / cesto inspecionado)",
      "Condições meteorológicas avaliadas",
      "Plano de resgate e pessoa apta presente",
      "Área abaixo isolada contra queda de galhos/ferramentas",
      "Motosserra não usada acima dos ombros / com uma mão"]
F6 = ["Freio de corrente presente e funcionando",
      "Pino pega-corrente",
      "Protetor da mão direita",
      "Protetor da mão esquerda",
      "Trava de segurança do acelerador",
      "Manual em português disponível",
      "Operador com treinamento de motosserra (mín. 8 h)",
      "Corrente e sabre em bom estado",
      "Abastecimento com motor desligado, longe de ignição",
      "Protetor do sabre no transporte"]
r = 2
for ficha, norma, items in (("F4", "NR-10", F4), ("F5", "NR-35", F5), ("F6", "NR-12 An. V", F6)):
    for i, it in enumerate(items, 1):
        ws.cell(row=r, column=1, value=ficha); ws.cell(row=r, column=2, value=norma)
        ws.cell(row=r, column=3, value=i); ws.cell(row=r, column=4, value=it)
        r += 1
last_v = r - 1
paint(ws, 2, last_v, 1, 4, CALC); paint(ws, 2, last_v, 5, 9, INPUT)
dvc = DataValidation(type="list", formula1='"C,NC,NA"', allow_blank=True)
ws.add_data_validation(dvc); dvc.add(f"G2:G{last_v}")
ws.conditional_formatting.add(f"G2:G{last_v}", CellIsRule(operator="equal", formula=['"NC"'], fill=red))
p_note = last_v + 2
ws.cell(row=p_note, column=4, value="Para registrar o mesmo item em outra data/equipamento, copie a linha para baixo.")

# ----------------------------------------------------------------- EPI ------
ws = wb.create_sheet("EPI")
cols = ["EPI", "Listado no PGR?", "Podador", "Ficha de entrega?", "Em uso no campo?",
        "CA válido?", "Nº da foto", "Estado / obs."]
header(ws, cols, [52, 12, 10, 13, 14, 11, 11, 34])
epis = [("Capacete classe B (CA 29792)", "Sim"), ("Óculos (CA 10346)", "Sim"),
        ("Protetor auditivo (CA 35981)", "Sim"),
        ("Luva contra agentes mecânicos (CA 40029)", "Sim"),
        ("Perneira (CA 25388)", "Sim"), ("Botina tipo B (CA 12160)", "Sim"),
        ("Cinturão paraquedista e talabarte (CA 35509)", "Sim"),
        ("Calça/perneira de proteção contra corte por motosserra", "Não"),
        ("Protetor facial (viseira de tela)", "Não"),
        ("Luva de proteção para motosserra", "Não"),
        ("Vestimenta para risco elétrico/arco", "Não"),
        ("Colete refletivo", "Não"), ("Protetor solar", "Medida adm.")]
r = 2
for pod in ("P1", "P2"):
    for nome, pgr in epis:
        ws.cell(row=r, column=1, value=nome); ws.cell(row=r, column=2, value=pgr)
        ws.cell(row=r, column=3, value=pod); r += 1
last_e = r - 1
paint(ws, 2, last_e, 1, 3, CALC); paint(ws, 2, last_e, 4, 8, INPUT)
dvs = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
dvu = DataValidation(type="list", formula1='"Sim,Não,Parcial"', allow_blank=True)
for d in (dvs, dvu):
    ws.add_data_validation(d)
dvs.add(f"D2:D{last_e}"); dvs.add(f"F2:F{last_e}"); dvu.add(f"E2:E{last_e}")
ws.conditional_formatting.add(f"E2:E{last_e}", CellIsRule(operator="equal", formula=['"Não"'], fill=red))

# --------------------------------------------------------------- Fotos ------
ws = wb.create_sheet("Fotos")
header(ws, ["Nº foto", "Data", "O que mostra", "Ficha/item", "Precisa desfocar?", "Arquivo"],
       [9, 11, 50, 12, 16, 28])
paint(ws, 2, 81, 1, 6, INPUT)
dvf = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
ws.add_data_validation(dvf); dvf.add("E2:E81")

# -------------------------------------------------------------- Resumo ------
ws = wb.create_sheet("Resumo")
ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 16
rows = [
    ("RUÍDO (amostras válidas)", None),
    ("Amostras lançadas", '=COUNTA(Dosimetria!L2:L13)'),
    ("Amostras válidas (calibração ±1 dB)", '=COUNTIF(Dosimetria!I2:I13,"Sim")'),
    ("NEN mínimo (dB(A))", '=IF(COUNTIF(Dosimetria!I2:I13,"Sim")=0,"",_xlfn.MINIFS(Dosimetria!P2:P13,Dosimetria!I2:I13,"Sim"))'),
    ("NEN máximo (dB(A))", '=IF(COUNTIF(Dosimetria!I2:I13,"Sim")=0,"",_xlfn.MAXIFS(Dosimetria!P2:P13,Dosimetria!I2:I13,"Sim"))'),
    ("Amostras acima do limite NHO-01 (NEN ≥ 85)", '=COUNTIF(Dosimetria!Q2:Q13,"Acima do limite")'),
    ("Amostras no nível de ação (82 ≤ NEN < 85)", '=COUNTIF(Dosimetria!Q2:Q13,"Nível de ação")'),
    ("Amostras acima do LT da NR-15 (dose > 100%)", '=COUNTIF(Dosimetria!R2:R13,"Acima do LT (insalubre)")'),
    ("", None),
    ("CONFORMIDADE (itens verificados)", None),
]
for norma in ("NR-10", "NR-35", "NR-12 An. V"):
    rows += [
        (f"{norma} — itens conformes (C)", f'=COUNTIFS(Verificacao!B:B,"{norma}",Verificacao!G:G,"C")'),
        (f"{norma} — itens não conformes (NC)", f'=COUNTIFS(Verificacao!B:B,"{norma}",Verificacao!G:G,"NC")'),
        (f"{norma} — % de conformidade (C / (C+NC))",
         f'=IFERROR(COUNTIFS(Verificacao!B:B,"{norma}",Verificacao!G:G,"C")/'
         f'(COUNTIFS(Verificacao!B:B,"{norma}",Verificacao!G:G,"C")+'
         f'COUNTIFS(Verificacao!B:B,"{norma}",Verificacao!G:G,"NC")),"")'),
    ]
rows += [("", None), ("EPI", None),
         ("EPIs não usados no campo (registros \"Não\")", '=COUNTIF(EPI!E:E,"Não")'),
         ("EPIs usados parcialmente", '=COUNTIF(EPI!E:E,"Parcial")'),
         ("EPIs de motosserra em uso (itens fora do PGR marcados \"Sim\")",
          '=COUNTIFS(EPI!B:B,"Não",EPI!E:E,"Sim")')]
for i, (lab, f) in enumerate(rows, 1):
    a = ws.cell(row=i, column=1, value=lab)
    if f is None and lab:
        a.font = Font(bold=True, color="16233F")
    elif f:
        c = ws.cell(row=i, column=2, value=f); c.fill = CALC; c.border = BOX
        if "% de conformidade" in lab:
            c.number_format = "0%"

# ---------------------------------------------------------- Referência ------
ws = wb.create_sheet("Referencia")
ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 30
ref = [("Parâmetro", "NHO-01 (canal 1)", "NR-15 Anexo 1 (canal 2)"),
       ("Critério de referência", "85 dB(A) / 8 h", "85 dB(A) / 8 h"),
       ("Incremento de duplicação de dose (q)", "3 dB", "5 dB"),
       ("Nível limiar de integração", "80 dB(A)", "80 dB(A)"),
       ("Ponderação / resposta", "A / lenta", "A / lenta"),
       ("Nível de ação", "NEN 82 dB(A) (dose 50%)", "—"),
       ("Limite", "NEN 85 dB(A) (dose 100%)", "dose 100%"),
       ("Validade da amostra", "calibração antes × depois ≤ ±1 dB", ""),
       ("", "", ""),
       ("NR-10 — raios de delimitação (m)", "Rr (zona de risco)", "Rc (zona controlada)"),
       ("< 1 kV", "0,20", "0,70"),
       ("≥ 10 kV e < 15 kV (ex.: 13,8 kV)", "0,38", "1,38"),
       ("≥ 15 kV e < 20 kV", "0,40", "1,40"),
       ("Conferir a tabela no texto oficial da NR-10 antes da redação.", "", "")]
for i, row in enumerate(ref, 1):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v); c.border = BOX if v else Border()
        if i in (1, 10):
            c.fill = HEAD; c.font = WHITE

wb.save(OUT)
print(f"OK: {OUT}")
