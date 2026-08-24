# -*- coding: utf-8 -*-
"""
Gerador do pôster do TCC (Psike) a partir do modelo oficial P_STER_MODELO_2024.pptx.

Preenche as caixas de texto do modelo (Introdução/Objetivo, Metodologia,
Estudo de Caso, Resultados e Discussão, Conclusão, Referências e o cabeçalho
tema/autor/e-mail) com o conteúdo da monografia, PRESERVANDO o layout, as
cores, as fontes e as logomarcas do modelo. Trechos [PREENCHER: ...] dependem
de dados reais de campo e ficam para o autor completar; as figuras de exemplo
do modelo devem ser trocadas por gráficos/telas reais no PowerPoint.

Uso:  python3 gerar_poster.py
Requer: pip install python-pptx
Entrada: _modelo_original.pptx (o modelo oficial, versionado ao lado)
Saída:   Poster_Psike_TCC.pptx  (exportar para PDF no PowerPoint para entrega)
"""
import copy
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = "_modelo_original.pptx"
OUT = "Poster_Psike_TCC.pptx"

# ---- Conteúdo por seção (título mantido do modelo; corpo substituído) --------

INTRO = [
    "A NR-1 exige que o empregador analise os acidentes e as doenças "
    "relacionadas ao trabalho, identifique suas causas e realimente o PGR "
    "(GRO). Na prática, o registro segue manual e tardio, a investigação para "
    "no \u201cato inseguro\u201d e os quase acidentes raramente são reportados.",
    "No setor elétrico a gravidade é extrema: 2.089 acidentes de origem "
    "elétrica no Brasil em 2023, com 781 óbitos (ABRACOPEL, 2024); 250 mortes "
    "envolvendo a rede de distribuição (ABRADEE, 2024). Cada ocorrência — e "
    "cada quase acidente — é aprendizado que não pode ser desperdiçado.",
    "A literatura brasileira consolidou o método da árvore de causas como "
    "antídoto à investigação culpabilizadora, revelando os fatores gerenciais "
    "e organizacionais na gênese dos acidentes (BINDER; ALMEIDA, 1997).",
    "Objetivo: desenvolver e aplicar um sistema digital de registro e "
    "investigação de acidentes e quase acidentes estruturado no método da "
    "árvore de causas, integrado ao GRO/PGR da NR-1, em uma organização de "
    "distribuição de energia elétrica.",
]

METODOLOGIA = [
    "Pesquisa aplicada, de natureza tecnológica: desenvolvimento de artefato "
    "(o sistema digital, módulo central da plataforma Psike) seguido de estudo "
    "de caso — análise documental retrospectiva de ocorrências anonimizadas e "
    "piloto de campo.",
    "Arquitetura de custo zero: frontend em página web autocontida (sem "
    "instalação, uso no celular em campo) e backend gratuito em nuvem como "
    "repositório único. Tratamento de dados anonimizado, conforme a LGPD.",
    "Registro em campo em menos de 5 minutos: classificação (acidente com/sem "
    "afastamento, quase acidente, condição insegura), descrição, tarefa e "
    "alerta automático de CAT. Investigação guiada em 3 níveis de causas: "
    "imediatas → subjacentes → básicas — o formulário não permite concluir só "
    "com causas imediatas.",
    "Ações corretivas/preventivas com responsável e prazo; indicadores "
    "automáticos: razão quase acidentes/acidentes, tempo evento–registro, "
    "ações no prazo e taxas de frequência e gravidade (NBR 14280). Módulos "
    "complementares na mesma base: permissão de trabalho (NR-10/NR-35), "
    "fatores psicossociais (NR-1) e AEP (NR-17).",
]

ESTUDO_CASO = [
    "Organização do setor de distribuição de energia elétrica, preservada em "
    "anonimato. Parte documental: [PREENCHER: número] ocorrências históricas "
    "anonimizadas reinvestigadas com o roteiro de 3 níveis, comparando a "
    "profundidade causal com a análise original.",
    "Parte de campo: piloto no período de [PREENCHER: período], com as equipes "
    "de [PREENCHER: escopo] registrando novas ocorrências e condições "
    "inseguras no local, pelo celular.",
    "Nenhum caso identifica pessoas, datas exatas ou locais; os exemplos são "
    "descritos com função genérica e circunstâncias descaracterizadas.",
]

RESULTADOS = [
    "O sistema foi implementado integralmente e está operacional: registro no "
    "local, investigação em 3 níveis, ações e indicadores atualizados sem "
    "transcrição manual, a custo de operação nulo.",
    "[PREENCHER: inserir resultados reais — distribuição de causas por nível "
    "(análise original × estruturada), volume de quase acidentes reportados no "
    "piloto, razão QA/acidentes e tempo evento–registro. Substituir as figuras "
    "pelo gráfico do painel e telas do módulo.]",
    "Discussão: quando o instrumento obriga a progressão até as causas "
    "básicas, a investigação alcança os fatores organizacionais que a análise "
    "tradicional não vê (BINDER; ALMEIDA, 1997; REASON, 1997) — e o reporte "
    "fácil de quase acidentes captura a base da pirâmide de eventos que "
    "antecede a lesão grave (HEINRICH, 1931; BIRD; GERMAIN, 1985).",
]

CONCLUSAO = [
    "A digitalização do ciclo registro–investigação–ação torna exequível, a "
    "custo praticamente nulo, a exigência da NR-1 de analisar ocorrências e "
    "realimentar o PGR — deslocando a análise da culpabilização individual "
    "para os fatores organizacionais e reduzindo a subnotificação de quase "
    "acidentes.",
    "Trabalhos futuros: diagrama completo da árvore de causas no software, "
    "autenticação e segregação por organização, acompanhamento longitudinal "
    "dos indicadores e cruzamento analítico entre ocorrências, permissões de "
    "trabalho e sinalização psicossocial.",
]

REFERENCIAS = [
    "ABRACOPEL. Anuário Estatístico de Acidentes de Origem Elétrica 2024 "
    "(ano-base 2023). Salto, 2024.",
    "BINDER, M. C. P.; ALMEIDA, I. M. Estudo de caso de dois acidentes do "
    "trabalho investigados com o método de árvore de causas. Cad. Saúde "
    "Pública, v. 13, n. 4, 1997.",
    "BRASIL. NR-1 — Gerenciamento de riscos ocupacionais (Portaria MTE nº "
    "1.419/2024).",
    "ABNT. NBR 14280:2001 — Cadastro de acidente do trabalho.",
    "REASON, J. Managing the risks of organizational accidents. Ashgate, 1997.",
    "HEINRICH, H. W. Industrial accident prevention. McGraw-Hill, 1931.",
]

HEADER = {
    "tema": "SISTEMA DIGITAL DE INVESTIGAÇÃO DE ACIDENTES PELA ÁRVORE DE "
            "CAUSAS: registro de ocorrências integrado à NR-1 na distribuição "
            "de energia elétrica",
    "autor": "Rodrigo Zambon [PREENCHER: nome completo]",
    "contato": "[PREENCHER: e-mail]  ·  Supervisão: [PREENCHER: nome]  ·  "
               "Apoio: CERPRO",
}

# ------------------------------------------------------------------ helpers ---

def _first_run(paragraph):
    runs = paragraph.findall(qn("a:r"))
    return runs[0] if runs else None

def _run_with_text(paragraph):
    """Primeiro run que possui elemento <a:t> (o que carrega texto de fato);
    cai para o primeiro run se nenhum tiver."""
    runs = paragraph.findall(qn("a:r"))
    for r in runs:
        if r.find(qn("a:t")) is not None:
            return r
    return runs[0] if runs else None

def replace_body(text_frame, paragraphs):
    """Mantém o 1º parágrafo (título da seção) e substitui o corpo,
    clonando o parágrafo de corpo do modelo para preservar fonte/cor/tamanho."""
    ps = text_frame.paragraphs
    if len(ps) < 2 or not paragraphs:
        return
    txBody = text_frame._txBody
    title_p = ps[0]._p
    # modelo de corpo = primeiro parágrafo (após o título) que tenha texto real
    tmpl_p = None
    for p in ps[1:]:
        if _run_with_text(p._p) is not None:
            tmpl_p = p._p
            break
    if tmpl_p is None:
        tmpl_p = ps[1]._p
    # remove todos os parágrafos exceto o título
    for p in list(txBody.findall(qn("a:p"))):
        if p is not title_p:
            txBody.remove(p)
    # recria o corpo a partir do template clonado
    for txt in paragraphs:
        newp = copy.deepcopy(tmpl_p)
        keep = _run_with_text(newp)
        for r in newp.findall(qn("a:r")):
            if r is not keep:
                newp.remove(r)
        if keep is not None:
            t = keep.find(qn("a:t"))
            if t is None:
                t = keep.makeelement(qn("a:t"), {})
                keep.append(t)
            t.text = txt
        txBody.append(newp)

def replace_lines(text_frame, lines):
    """Substitui texto linha a linha (cabeçalho), preservando formatação."""
    ps = text_frame.paragraphs
    for i, line in enumerate(lines):
        if i >= len(ps):
            break
        r = _first_run(ps[i]._p)
        if r is not None:
            t = r.find(qn("a:t"))
            if t is not None:
                t.text = line
        # remove runs extras da linha para não sobrar lorem
        for extra in ps[i]._p.findall(qn("a:r"))[1:]:
            ps[i]._p.remove(extra)

# --------------------------------------------------------------------- main ---

prs = Presentation(SRC)
slide = prs.slides[0]

by_name = {sh.name: sh for sh in slide.shapes if sh.has_text_frame}

SECTIONS = {
    "CaixaDeTexto 25": INTRO,          # Introdução/Objetivo
    "CaixaDeTexto 29": METODOLOGIA,    # Metodologia
    "CaixaDeTexto 27": ESTUDO_CASO,    # Estudo de Caso
    "CaixaDeTexto 44": RESULTADOS,     # Resultados e Discussão
    "CaixaDeTexto 46": CONCLUSAO,      # Conclusão
    "CaixaDeTexto 47": REFERENCIAS,    # Referências
}
for name, content in SECTIONS.items():
    if name in by_name:
        replace_body(by_name[name].text_frame, content)

# Cabeçalho (TEMA / autor / contato)
if "CaixaDeTexto 2" in by_name:
    replace_lines(by_name["CaixaDeTexto 2"].text_frame,
                  [HEADER["tema"], HEADER["autor"], HEADER["contato"]])

prs.save(OUT)
print(f"OK: {OUT} gerado a partir de {SRC}.")
