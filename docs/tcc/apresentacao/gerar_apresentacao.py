# -*- coding: utf-8 -*-
"""
Gerador da apresentação da banca (defesa) do TCC — Psike.
Produz Apresentacao_Banca_Psike.pptx (16:9), com a identidade visual do
projeto (navy #16233F + dourado #B4842A) e telas reais da plataforma.

Uso:  python3 gerar_apresentacao.py
Requer: pip install python-pptx  (e as imagens em ./img)
Notas para a defesa (roteiro de fala) ficam nas notas de cada slide.
Trechos [PREENCHER: ...] dependem de dados reais e devem ser completados.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY   = RGBColor(0x16, 0x23, 0x3F)
NAVY2  = RGBColor(0x1F, 0x31, 0x57)
GOLD   = RGBColor(0xB4, 0x84, 0x2A)
PAPER  = RGBColor(0xF1, 0xF3, 0xF6)
SLATE  = RGBColor(0x5B, 0x65, 0x76)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LINE   = RGBColor(0xDD, 0xE2, 0xE9)

IMG = os.path.join(os.path.dirname(__file__), "img")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def box(s, l, t, w, h):
    tb = s.shapes.add_textbox(l, t, w, h)
    tb.text_frame.word_wrap = True
    return tb

def rect(s, l, t, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def set_para(p, text, size, color, bold=False, align=PP_ALIGN.LEFT,
             font="Calibri", space_after=8):
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    for r in p.runs:
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold; r.font.name = font
    return p

def title_bar(s, eyebrow, title):
    """Barra de título padrão dos slides de conteúdo."""
    rect(s, 0, 0, SW, Inches(0.14), GOLD)
    tb = box(s, Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.15))
    tf = tb.text_frame
    set_para(tf.paragraphs[0], eyebrow.upper(), 12, GOLD, bold=True,
             font="Consolas", space_after=2)
    set_para(tf.add_paragraph(), title, 30, NAVY, bold=True, font="Georgia")

def bullets(s, items, top=Inches(1.7), left=Inches(0.7),
            width=Inches(12.0), height=Inches(5.2), size=18, gap=10):
    tb = box(s, left, top, width, height)
    tf = tb.text_frame
    first = True
    for it in items:
        lvl = 0
        txt = it
        if isinstance(it, tuple):
            txt, lvl = it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        bullet = "•  " if lvl == 0 else "–  "
        set_para(p, bullet + txt, size - (2 if lvl else 0),
                 NAVY if lvl == 0 else SLATE, space_after=gap)
        p.level = lvl
    return tb

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

def pic_card(s, path, l, t, w):
    """Insere imagem com borda, ajustando altura pela proporção."""
    if not os.path.exists(path):
        ph = rect(s, l, t, w, Inches(3.5), PAPER)
        set_para(ph.text_frame.paragraphs[0], "[imagem]", 14, SLATE,
                 align=PP_ALIGN.CENTER)
        return ph
    from PIL import Image
    iw, ih = Image.open(path).size
    h = int(w * ih / iw)
    pic = s.shapes.add_picture(path, l, t, width=w, height=h)
    pic.line.color.rgb = LINE; pic.line.width = Pt(1)
    return pic

# ============================================================ 1. CAPA =========
s = slide(); bg(s, NAVY)
rect(s, 0, Inches(6.9), SW, Inches(0.14), GOLD)
mk = rect(s, Inches(0.7), Inches(0.7), Inches(0.7), Inches(0.7), GOLD)
set_para(mk.text_frame.paragraphs[0], "P", 26, NAVY, bold=True,
         align=PP_ALIGN.CENTER, font="Georgia")
mk.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
tb = box(s, Inches(1.6), Inches(0.72), Inches(6), Inches(0.7))
set_para(tb.text_frame.paragraphs[0], "Psike", 26, WHITE, bold=True, font="Georgia")

tb = box(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.9))
tf = tb.text_frame
set_para(tf.paragraphs[0], "INVESTIGAÇÃO DIGITAL DE ACIDENTES PELO MÉTODO DA "
         "ÁRVORE DE CAUSAS", 32, WHITE, bold=True, font="Georgia", space_after=10)
set_para(tf.add_paragraph(),
         "Registro e análise de acidentes e quase acidentes integrados à NR-1 "
         "em serviços de distribuição de energia elétrica", 20,
         RGBColor(0xC7,0xCE,0xDB), font="Calibri")

tb = box(s, Inches(0.9), Inches(5.2), Inches(11.5), Inches(1.6))
tf = tb.text_frame
set_para(tf.paragraphs[0], "Rodrigo Zambon  [PREENCHER: nome completo]",
         18, GOLD, bold=True, space_after=4)
set_para(tf.add_paragraph(),
         "Supervisão: [PREENCHER]   ·   Especialização em Engenharia de "
         "Segurança do Trabalho — PECE/EPUSP", 14, RGBColor(0xAE,0xB8,0xC9))
set_para(tf.add_paragraph(), "São Paulo · [PREENCHER: data da defesa]",
         14, RGBColor(0xAE,0xB8,0xC9))
notes(s, "Cumprimentar a banca. Frase de abertura: no setor elétrico o acidente "
         "raramente é leve — e a NR-1 obriga a analisar cada ocorrência e "
         "realimentar o PGR. Este TCC entrega o instrumento que torna isso "
         "exequível. Duração-alvo: ~15 min + arguição.")

# ============================================ 2. CONTEXTO / PROBLEMA ===========
s = slide(); bg(s, WHITE)
title_bar(s, "Contexto e problema", "A norma exige analisar; a prática ainda é manual")
bullets(s, [
    "A NR-1 (GRO) obriga o empregador a analisar acidentes e doenças "
    "relacionadas ao trabalho, identificar causas e realimentar o PGR.",
    "Na prática: registro em papel, horas/dias após o evento; investigação "
    "que para no “ato inseguro”; quase acidentes raramente reportados; ações "
    "sem acompanhamento.",
    "A literatura chama essa prática de paradigma culpabilizador — e a "
    "associa diretamente à recorrência dos eventos (Binder; Almeida, 1997; "
    "Almeida, 2006).",
    ("Setor elétrico: 2.089 acidentes de origem elétrica em 2023, 781 óbitos "
     "(Abracopel, 2024); 250 mortes na rede de distribuição (Abradee, 2024).", 1),
    ("Eventos raros e graves ⇒ o quase acidente é o insumo mais abundante do "
     "aprendizado — se for reportado.", 1),
])
notes(s, "Ancorar o problema nos números do setor e na exigência da NR-1. "
         "Ponto-chave: investigação rasa = acidente repetido.")

# ============================================ 3. JUSTIFICATIVA =================
s = slide(); bg(s, WHITE)
title_bar(s, "Justificativa", "Três razões para este trabalho")
bullets(s, [
    "Gravidade setorial: letalidade dos acidentes elétricos supera 1/3 das "
    "ocorrências (Abracopel, 2024) — cada evento desperdiçado custa caro.",
    "Exigência normativa: NR-1 demanda análise com método e registro "
    "rastreável; NBR 14280 padroniza cadastro e estatísticas.",
    "Lacuna prática: o método da árvore de causas é consolidado na literatura "
    "brasileira desde os anos 1990, mas carece de ferramenta digital "
    "acessível que o leve ao campo.",
    "Proposta: sistema de custo praticamente nulo, replicável por equipes de "
    "SST enxutas, sem apoio de TI.",
])
notes(s, "Deixar claro o gap: método existe, norma exige, ferramenta "
         "acessível não existia. O diferencial é operacionalizar.")

# ============================================ 4. OBJETIVOS =====================
s = slide(); bg(s, WHITE)
title_bar(s, "Objetivos", "Objetivo geral e específicos")
bullets(s, [
    "GERAL: desenvolver e aplicar um sistema digital de registro e "
    "investigação de acidentes e quase acidentes, estruturado no método da "
    "árvore de causas e integrado ao GRO (NR-1), em uma organização de "
    "distribuição de energia elétrica.",
    ("Sistematizar os requisitos da NR-1, NBR 14280, ISO 45001 e a "
     "articulação com a CAT/eSocial;", 1),
    ("Revisar os modelos de causalidade e o método da árvore de causas;", 1),
    ("Implementar registro em campo, investigação em 3 níveis de causas, "
     "gestão de ações e indicadores;", 1),
    ("Integrar aos módulos complementares (PT NR-10/35, psicossocial NR-1, "
     "AEP NR-17) na mesma base;", 1),
    ("Aplicar em estudo de caso (retrospectivo + piloto) e discutir "
     "resultados e limitações.", 1),
])
notes(s, "Ler o geral com calma; os específicos espelham a estrutura do "
         "trabalho.")

# ============================================ 5. FUNDAMENTAÇÃO =================
s = slide(); bg(s, WHITE)
title_bar(s, "Fundamentação", "Do modelo causal ao roteiro de investigação")
bullets(s, [
    "Heinrich (1931) e Bird & Germain (1985): a pirâmide de eventos — muitos "
    "quase acidentes antecedem a lesão grave; reportá-los é prevenção.",
    "Reason (1990; 1997): falhas ativas × condições latentes — investigar só "
    "a conduta do trabalhador deixa intactas as causas organizacionais.",
    "Árvore de causas (INRS; Binder; Monteau; Almeida): reconstrução da rede "
    "de fatos, sem juízo de culpa — referência brasileira consolidada.",
    "Fatores humanos e organizacionais nas causas básicas: fadiga, pressão "
    "de tempo, sobrecarga — ponte com os riscos psicossociais da NR-1 e com "
    "o burnout (CID-11/QD85; Treml et al., 2025).",
    "Operacionalização adotada: 3 níveis obrigatórios — causas imediatas → "
    "subjacentes → básicas; o formulário não conclui no nível 1.",
], size=17)
notes(s, "Mostrar a linha: modelo causal define a investigação. Citar o "
         "burnout como fator humano contribuinte (ponto cobrado na "
         "avaliação). Fechar com a tradução do método para o software.")

# ============================================ 6. METODOLOGIA ===================
s = slide(); bg(s, WHITE)
title_bar(s, "Metodologia", "Artefato + estudo de caso em duas partes")
bullets(s, [
    "Pesquisa aplicada/tecnológica: desenvolvimento de artefato + estudo de "
    "caso em organização do setor elétrico (anonimizada).",
    "Parte documental: ocorrências históricas anonimizadas reinvestigadas com "
    "o roteiro de 3 níveis — comparação com a análise original.",
    "Parte de campo: piloto com as equipes registrando novas ocorrências e "
    "condições inseguras no local, pelo celular.",
    "Arquitetura: custo zero · sem instalação · página web autocontida + "
    "backend gratuito em nuvem; dados tratados de forma agregada (LGPD).",
], height=Inches(3.4))
notes(s, "Destacar o desenho em duas partes: o retrospectivo compara "
         "profundidade causal; o piloto mede reporte e tempo de registro. "
         "Anonimato: sem nomes, datas exatas ou locais.")

# ============================================ 7. O SISTEMA (MÓDULOS) ===========
s = slide(); bg(s, PAPER)
title_bar(s, "O sistema", "Módulo central + três módulos integrados")
mods = [
    ("CENTRAL", "Acidentes e quase acidentes", "NR-1 · NBR 14280 · árvore de causas"),
    ("MÓDULO 2", "Permissão de trabalho", "NR-10 · NR-35 (distribuição)"),
    ("MÓDULO 3", "Riscos psicossociais", "NR-1 · fatores humanos"),
    ("MÓDULO 4", "Ergonomia (AEP)", "NR-17"),
]
cw, gap = Inches(2.95), Inches(0.25)
x0 = Inches(0.7); y0 = Inches(2.1)
for i, (num, tit, sub) in enumerate(mods):
    x = x0 + i * (cw + gap)
    card = rect(s, x, y0, cw, Inches(3.4), WHITE)
    tf = box(s, x + Inches(0.2), y0 + Inches(0.25), cw - Inches(0.4), Inches(3.0)).text_frame
    set_para(tf.paragraphs[0], "● em uso", 11, RGBColor(0x2E,0x7D,0x63),
             bold=True, font="Consolas", space_after=8)
    set_para(tf.add_paragraph(), num, 11, GOLD, bold=True, font="Consolas",
             space_after=2)
    set_para(tf.add_paragraph(), tit, 17, NAVY, bold=True, font="Georgia",
             space_after=6)
    set_para(tf.add_paragraph(), sub, 12, SLATE)
notes(s, "O módulo de acidentes é o coração do TCC; os demais alimentam a "
         "mesma base — dá para cruzar ocorrência × permissão emitida × "
         "sinalização psicossocial do setor.")

# ============================================ 8. MÓDULO CENTRAL — FLUXO ========
s = slide(); bg(s, WHITE)
title_bar(s, "O módulo central em detalhe",
          "Registro no local → 3 níveis de causas → ações e indicadores")
pic_card(s, os.path.join(IMG, "shot_acidentes.png"), Inches(0.7), Inches(1.9), Inches(6.0))
tb = box(s, Inches(7.1), Inches(1.9), Inches(5.6), Inches(4.6))
tf = tb.text_frame; tf.word_wrap = True
set_para(tf.paragraphs[0], "Fluxo da investigação", 16, NAVY, bold=True,
         font="Georgia", space_after=8)
for txt in [
    "1. Registro em < 5 min: tipo, tarefa, descrição — alerta de CAT "
    "automático quando aplicável.",
    "2. Causas imediatas: o que produziu o dano na cena.",
    "3. Causas subjacentes: o que na tarefa/posto tornou possível.",
    "4. Causas básicas: o que na gestão/organização originou — inclui "
    "fatores humanos (fadiga, sobrecarga).",
    "5. Ações com responsável, prazo e status; indicadores automáticos "
    "(razão QA/acidentes, TF/TG da NBR 14280).",
]:
    set_para(tf.add_paragraph(), txt, 13, SLATE, space_after=6)
notes(s, "Ponto de venda técnico: o formulário NÃO permite concluir só com "
         "causas imediatas — é o método da árvore traduzido em software. "
         "Se possível, demo ao vivo aqui.")

# ============================================ 9. MÓDULOS COMPLEMENTARES ========
s = slide(); bg(s, WHITE)
title_bar(s, "Módulos integrados", "Permissão de trabalho, psicossocial e ergonomia")
pic_card(s, os.path.join(IMG, "shot_apr.png"), Inches(0.6), Inches(1.9), Inches(4.0))
pic_card(s, os.path.join(IMG, "shot_form.png"), Inches(4.75), Inches(1.9), Inches(4.0))
pic_card(s, os.path.join(IMG, "shot_aep.png"), Inches(8.9), Inches(1.9), Inches(3.7))
tb = box(s, Inches(0.6), Inches(6.5), Inches(12), Inches(0.7))
set_para(tb.text_frame.paragraphs[0],
         "PT: checklist NR-10/35, geolocalização e assinatura  ·  "
         "Psicossocial: 10 itens anônimos, 6 dimensões  ·  AEP: 5 blocos com "
         "parecer", 13, SLATE, align=PP_ALIGN.CENTER)
notes(s, "Passar rápido — a mensagem é integração: a mesma base de dados "
         "permite cruzar ocorrências com permissões e com a sinalização "
         "psicossocial (fatores humanos).")

# ============================================ 10. RESULTADOS ===================
s = slide(); bg(s, WHITE)
title_bar(s, "Resultados", "Sistema operacional + estudo de caso")
bullets(s, [
    "Sistema implementado integralmente: registro, investigação em 3 níveis, "
    "ações e indicadores — sem transcrição manual, custo de operação nulo.",
    "[PREENCHER: parte documental — distribuição de causas por nível na "
    "análise original × estruturada; 2–3 casos anonimizados].",
    "[PREENCHER: piloto — volume por tipo, razão quase acidentes/acidentes, "
    "tempo mediano evento–registro, ações no prazo; inserir gráficos].",
    "[PREENCHER: revisões do IRO/plano de ação decorrentes, sem identificar "
    "a organização].",
], height=Inches(3.6))
notes(s, "Dados reais aqui — NÃO inventar. Se o piloto estiver em curso na "
         "defesa, apresentar a parte documental como resultado principal e o "
         "piloto como em andamento.")

# ============================================ 11. DISCUSSÃO / LIMITAÇÕES =======
s = slide(); bg(s, WHITE)
title_bar(s, "Discussão e limitações", "O que os resultados indicam — e o que não")
bullets(s, [
    "Se a investigação estruturada revelou mais causas básicas: o instrumento "
    "condiciona a profundidade — confirma Binder & Almeida (1997).",
    "Se o reporte de quase acidentes cresceu: a subnotificação era atrito de "
    "processo — o sistema captura a base da pirâmide (Heinrich; Bird).",
    "Fatores humanos entre as causas básicas articulam acidentes e gestão "
    "psicossocial da NR-1 — faces do mesmo sistema.",
    ("Limitações: roteiro em 3 níveis é simplificação da árvore completa;", 1),
    ("qualidade dos registros históricos; piloto curto em uma organização;", 1),
    ("sem autenticação/segregação por empresa; limites da camada gratuita.", 1),
])
notes(s, "Reconhecer limites com maturidade — a banca valoriza. O roteiro de "
         "3 níveis é operacionalização, não o diagrama completo do método.")

# ============================================ 12. CONCLUSÃO ====================
s = slide(); bg(s, WHITE)
title_bar(s, "Conclusão", "Contribuição e trabalhos futuros")
bullets(s, [
    "A digitalização do ciclo registro–investigação–ação torna exequível a "
    "exigência da NR-1 de analisar ocorrências e realimentar o PGR — a custo "
    "praticamente nulo.",
    "Deslocamento essencial: da culpabilização individual para os fatores "
    "organizacionais — com rastreabilidade de auditoria.",
    "Trabalhos futuros: diagrama completo da árvore no software; autenticação "
    "e segregação por organização; leitura longitudinal dos indicadores; "
    "cruzamento ocorrências × permissões × psicossocial.",
])
notes(s, "Fechar reafirmando: método consolidado + ferramenta acessível = "
         "prevenção que aprende com cada evento. Gancho de continuidade.")

# ============================================ 13. ENCERRAMENTO =================
s = slide(); bg(s, NAVY)
rect(s, 0, Inches(6.9), SW, Inches(0.14), GOLD)
tb = box(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(2.2))
tf = tb.text_frame
set_para(tf.paragraphs[0], "Obrigado.", 40, WHITE, bold=True, font="Georgia",
         space_after=12)
set_para(tf.add_paragraph(),
         "Rodrigo Zambon  ·  [PREENCHER: e-mail]", 18, GOLD, bold=True,
         space_after=4)
set_para(tf.add_paragraph(),
         "Psike — investigação digital de acidentes pela árvore de causas (NR-1)",
         15, RGBColor(0xC7,0xCE,0xDB), space_after=10)
set_para(tf.add_paragraph(), "Agradecimento: CERPRO", 13,
         RGBColor(0xAE,0xB8,0xC9))
notes(s, "Agradecer à banca, à supervisão da monografia e à CERPRO. Ter o "
         "app aberto para eventual demonstração ao vivo.")

OUT = os.path.join(os.path.dirname(__file__), "Apresentacao_Banca_Psike.pptx")
prs.save(OUT)
print(f"OK: {os.path.basename(OUT)} gerado — {len(prs.slides.__iter__.__self__._sldIdLst)} slides.")
