# -*- coding: utf-8 -*-
"""
Gera Planilha_Dados_Fibra.xlsx para o levantamento de postes ocupados por
cabos de telecomunicações. Compara cada medição com o limite da aba
Referencia (editável) e calcula os totais do capítulo de resultados.

Uso:  python3 gerar_planilha_fibra.py      Requer: pip install openpyxl
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = "Planilha_Dados_Fibra.xlsx"
HEAD = PatternFill("solid", fgColor="16233F")
INPUT = PatternFill("solid", fgColor="FFF8DC")
CALC = PatternFill("solid", fgColor="E8EEF6")
RED = PatternFill("solid", fgColor="F4C7C3")
WHITE = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="B8C2D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
N = 300                      # linhas de postes
LAST = N + 1

wb = Workbook()


def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.fill = HEAD; cell.font = WHITE; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 54
    ws.freeze_panes = "B2"


def paint(ws, r1, r2, c1, c2, fill):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill; cell.border = BOX


def listv(ws, options, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(rng)


# ------------------------------------------------------------ Leia-me -------
ws = wb.active; ws.title = "Leia-me"
ws.column_dimensions["A"].width = 112
for i, t in enumerate([
    "PLANILHA DE DADOS — OCUPAÇÃO DE POSTES POR CABOS DE TELECOMUNICAÇÕES",
    "",
    "Células AMARELAS = você preenche.  Células AZUIS = calculadas (não editar).",
    "Sem nomes: postes T1-01, T1-02…; trechos T1…T4; provedores A, B, C…",
    "",
    "Postes      — uma linha por poste (ficha F2). Altura: preencha a medida direta (col. H)",
    "              OU distância horizontal + ângulo + altura dos olhos (cols. I-K).",
    "Equipes     — observações de equipes de telecom (ficha F3).",
    "Fotos       — registro fotográfico (ficha F5).",
    "Resumo      — totais e percentuais automáticos.",
    "Referencia  — limites usados nas comparações. CONFIRA na NBR 15214 ou na norma da",
    "              distribuidora e ajuste se necessário: os cálculos se atualizam.",
    "",
    "A altura que conta é a do ponto mais baixo do cabo, normalmente o meio do vão.",
], 1):
    ws.cell(row=i, column=1, value=t)
ws["A1"].font = Font(bold=True, size=13, color="16233F")

# ---------------------------------------------------------- Referência -----
ref = wb.create_sheet("Referencia")
ref.column_dimensions["A"].width = 44; ref.column_dimensions["B"].width = 14
ref.column_dimensions["C"].width = 60
rows = [("Local do vão", "Altura mín. (m)", "Fonte / observação"),
        ("Rua/avenida", 5.00, "NBR 15214 — conferir"),
        ("Entrada de veículos", 4.50, "NBR 15214 (entradas de prédios / uso restrito a veículos) — conferir"),
        ("Rodovia/ferrovia", 6.00, "6,00 a 7,00 m conforme a fonte (7,00 em rodovias federais em algumas normas) — conferir"),
        ("Só pedestres", None, "Preencher com o valor da norma da distribuidora, se houver"),
        ("", None, ""),
        ("Outros requisitos", "Valor", ""),
        ("Afastamento mín. telecom – neutro (m)", 0.60, "NBR 15214 — conferir"),
        ("Distância mín. da caixa de emenda/reserva ao poste (m)", 2.00, "NBR 15214 — registrado na col. U (S/N)"),
        ("Diâmetro máx. por ponto de fixação (mm)", 65, "NBR 15214 — referência para a discussão")]
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row, 1):
        c = ref.cell(row=i, column=j, value=v)
        if i in (1, 7):
            c.fill = HEAD; c.font = WHITE
        elif v is not None and v != "":
            c.border = BOX
for r in (2, 3, 4, 5, 8, 9, 10):
    ref.cell(row=r, column=2).fill = INPUT
AFAST = "Referencia!$B$8"

# ------------------------------------------------------------- Postes -------
ws = wb.create_sheet("Postes", 1)
cols = ["Poste", "Trecho", "Data", "Local do vão", "Estrutura elétrica",
        "Nº cabos telecom", "Nº pontos de fixação",
        "Altura mín. meio do vão — medida direta (m)", "Dist. horizontal (m) [método ângulo]",
        "Ângulo de visada (°)", "Altura dos olhos (m)", "Altura considerada (m)",
        "Limite (m)", "Altura OK?", "Altura do neutro no poste (m)",
        "Altura telecom mais alto no poste (m)", "Afastamento telecom–neutro (m)",
        "Afastamento OK?", "Plaqueta em todos os cabos?", "Sobra enrolada no poste?",
        "Caixa de emenda < 2 m do poste?", "Equipamento ativo no poste?",
        "Cabo solto/caído/abandonado?", "Cabo encostando na rede/ferragem?",
        "Poste danificado/inclinado?", "Obstrução da subida", "Nº de não conformidades",
        "Poste irregular?", "Nº das fotos", "Observações"]
header(ws, cols, [9, 7, 11, 17, 15, 9, 9, 13, 12, 10, 9, 11, 9, 9, 11, 11, 12,
                  11, 12, 11, 12, 11, 12, 12, 11, 11, 11, 10, 10, 28])
for r in range(2, LAST + 1):
    f = {
        12: f'=IF(H{r}<>"",H{r},IF(AND(I{r}<>"",J{r}<>""),ROUND(I{r}*TAN(RADIANS(J{r}))+IF(K{r}="",1.6,K{r}),2),""))',
        13: f'=IF(D{r}="","",IFERROR(IF(VLOOKUP(D{r},Referencia!$A$2:$B$5,2,FALSE)="","",VLOOKUP(D{r},Referencia!$A$2:$B$5,2,FALSE)),""))',
        14: f'=IF(OR(L{r}="",M{r}=""),"",IF(L{r}>=M{r},"C","NC"))',
        17: f'=IF(OR(O{r}="",P{r}=""),"",ROUND(O{r}-P{r},2))',
        18: f'=IF(Q{r}="","",IF(Q{r}>={AFAST},"C","NC"))',
        27: (f'=IF(F{r}="","",(N{r}="NC")+(R{r}="NC")+(S{r}="Não")+(S{r}="Parcial")'
             f'+(T{r}="Sim")+(U{r}="Sim")+(V{r}="Sim")+(W{r}="Sim")+(X{r}="Sim"))'),
        28: f'=IF(AA{r}="","",IF(AA{r}>0,"Sim","Não"))',
    }
    for c, v in f.items():
        ws.cell(row=r, column=c, value=v)
    ws.cell(row=r, column=3).number_format = "DD/MM/YYYY"
paint(ws, 2, LAST, 1, 11, INPUT); paint(ws, 2, LAST, 15, 16, INPUT)
paint(ws, 2, LAST, 19, 26, INPUT); paint(ws, 2, LAST, 29, 30, INPUT)
paint(ws, 2, LAST, 12, 14, CALC); paint(ws, 2, LAST, 17, 18, CALC)
paint(ws, 2, LAST, 27, 28, CALC)
listv(ws, ["T1", "T2", "T3", "T4", "T5", "T6"], f"B2:B{LAST}")
listv(ws, ["Rua/avenida", "Entrada de veículos", "Rodovia/ferrovia", "Só pedestres"], f"D2:D{LAST}")
listv(ws, ["Prim.+sec.", "Só primário", "Só secundário", "Com trafo/chave"], f"E2:E{LAST}")
listv(ws, ["Sim", "Não", "Parcial"], f"S2:S{LAST}")
for col in "TUVWXY":
    listv(ws, ["Sim", "Não"], f"{col}2:{col}{LAST}")
listv(ws, ["Baixa", "Média", "Alta"], f"Z2:Z{LAST}")
for col in ("N", "R"):
    ws.conditional_formatting.add(f"{col}2:{col}{LAST}",
                                  FormulaRule(formula=[f'{col}2="NC"'], fill=RED))
ws.conditional_formatting.add(f"AB2:AB{LAST}", FormulaRule(formula=['AB2="Sim"'], fill=RED))

# ------------------------------------------------------------- Equipes ------
ws = wb.create_sheet("Equipes")
items = ["Sinalização viária", "Escada apoiada no poste (não nos cabos) e estabilizada",
         "Escada isolante", "Conectado acima de 2 m", "Pessoa no solo apoiando/controlando trânsito",
         "Distância da rede elétrica", "Capacete com jugular, luvas, calçado",
         "Suspensão em chuva/vento", "Cabo novo na faixa e com plaqueta"]
header(ws, ["Data", "Trecho", "Provedor (letra)", "Nº pessoas", "Atividade"] + items
       + ["Nº de NC", "Obs."], [11, 7, 9, 8, 18] + [13] * 9 + [8, 30])
for r in range(2, 42):
    ws.cell(row=r, column=15, value=f'=IF(A{r}="","",COUNTIF(F{r}:N{r},"NC"))')
paint(ws, 2, 41, 1, 14, INPUT); paint(ws, 2, 41, 16, 16, INPUT); paint(ws, 2, 41, 15, 15, CALC)
listv(ws, ["C", "NC", "NA"], "F2:N41")
ws.conditional_formatting.add("F2:N41", FormulaRule(formula=['F2="NC"'], fill=RED))

# --------------------------------------------------------------- Fotos ------
ws = wb.create_sheet("Fotos")
header(ws, ["Nº foto", "Data", "Poste/trecho", "O que mostra", "Desfocar?", "Arquivo"],
       [9, 11, 12, 50, 11, 28])
paint(ws, 2, 151, 1, 6, INPUT)
listv(ws, ["Sim", "Não"], "E2:E151")

# -------------------------------------------------------------- Resumo ------
ws = wb.create_sheet("Resumo", 2)
ws.column_dimensions["A"].width = 58
for c in "BCDEFG":
    ws.column_dimensions[c].width = 12
P = f"Postes!$A$2:$AB${LAST}"
def col(c):
    return f"Postes!${c}$2:${c}${LAST}"
n_postes = f'COUNT({col("F")})'
lines = [
    ("GERAL", None),
    ("Postes levantados", f'={n_postes}'),
    ("Postes com pelo menos 1 não conformidade", f'=COUNTIF({col("AB")},"Sim")'),
    ("% de postes irregulares", f'=IFERROR(COUNTIF({col("AB")},"Sim")/{n_postes},"")'),
    ("Média de cabos de telecom por poste", f'=IFERROR(ROUND(AVERAGE({col("F")}),1),"")'),
    ("Máximo de cabos em um poste", f'=IFERROR(MAX({col("F")}),"")'),
    ("Média de pontos de fixação por poste", f'=IFERROR(ROUND(AVERAGE({col("G")}),1),"")'),
    ("", None),
    ("ALTURA AO SOLO (vãos)", None),
    ("Vãos avaliados", f'=COUNTIF({col("N")},"C")+COUNTIF({col("N")},"NC")'),
    ("Vãos abaixo da altura mínima", f'=COUNTIF({col("N")},"NC")'),
    ("% de vãos abaixo da altura mínima",
     f'=IFERROR(COUNTIF({col("N")},"NC")/(COUNTIF({col("N")},"C")+COUNTIF({col("N")},"NC")),"")'),
    ("Menor altura encontrada (m)", f'=IFERROR(IF(COUNT({col("L")})=0,"",MIN({col("L")})),"")'),
    ("", None),
    ("AFASTAMENTO TELECOM – REDE ELÉTRICA", None),
    ("Postes avaliados", f'=COUNTIF({col("R")},"C")+COUNTIF({col("R")},"NC")'),
    ("Postes com afastamento abaixo do mínimo", f'=COUNTIF({col("R")},"NC")'),
    ("% abaixo do mínimo",
     f'=IFERROR(COUNTIF({col("R")},"NC")/(COUNTIF({col("R")},"C")+COUNTIF({col("R")},"NC")),"")'),
    ("Menor afastamento encontrado (m)", f'=IFERROR(IF(COUNT({col("Q")})=0,"",MIN({col("Q")})),"")'),
    ("", None),
    ("FREQUÊNCIA DAS IRREGULARIDADES (nº de postes)", None),
    ("Sem plaqueta (não ou parcial)", f'=COUNTIF({col("S")},"Não")+COUNTIF({col("S")},"Parcial")'),
    ("Sobra de cabo enrolada no poste", f'=COUNTIF({col("T")},"Sim")'),
    ("Caixa de emenda a menos de 2 m do poste", f'=COUNTIF({col("U")},"Sim")'),
    ("Equipamento ativo no poste", f'=COUNTIF({col("V")},"Sim")'),
    ("Cabo solto, caído ou abandonado", f'=COUNTIF({col("W")},"Sim")'),
    ("Cabo encostando na rede ou ferragem", f'=COUNTIF({col("X")},"Sim")'),
    ("Poste danificado ou inclinado", f'=COUNTIF({col("Y")},"Sim")'),
    ("", None),
    ("OBSTRUÇÃO DA SUBIDA (risco ao eletricista)", None),
    ("Postes com obstrução alta", f'=COUNTIF({col("Z")},"Alta")'),
    ("Média de cabos nos postes com obstrução alta",
     f'=IFERROR(ROUND(AVERAGEIF({col("Z")},"Alta",{col("F")}),1),"")'),
    ("Média de cabos nos postes com obstrução baixa",
     f'=IFERROR(ROUND(AVERAGEIF({col("Z")},"Baixa",{col("F")}),1),"")'),
    ("", None),
    ("EQUIPES DE TELECOM OBSERVADAS", None),
    ("Equipes observadas", '=COUNTA(Equipes!A2:A41)'),
    ("Sem sinalização viária (NC)", '=COUNTIF(Equipes!F2:F41,"NC")'),
    ("Escada apoiada nos cabos / sem estabilização (NC)", '=COUNTIF(Equipes!G2:G41,"NC")'),
    ("Sem conexão acima de 2 m (NC)", '=COUNTIF(Equipes!I2:I41,"NC")'),
]
for i, (lab, f) in enumerate(lines, 1):
    a = ws.cell(row=i, column=1, value=lab)
    if f is None and lab:
        a.font = Font(bold=True, color="16233F")
    elif f:
        c = ws.cell(row=i, column=2, value=f); c.fill = CALC; c.border = BOX
        if lab.startswith("%"):
            c.number_format = "0.0%"
# por trecho
start = len(lines) + 2
ws.cell(row=start, column=1, value="POR TRECHO").font = Font(bold=True, color="16233F")
heads = ["Trecho", "Postes", "Irregulares", "% irregulares", "Vãos NC altura", "Afast. NC"]
for j, h in enumerate(heads, 1):
    c = ws.cell(row=start + 1, column=j, value=h); c.fill = HEAD; c.font = WHITE
for k, t in enumerate(["T1", "T2", "T3", "T4", "T5", "T6"]):
    r = start + 2 + k
    ws.cell(row=r, column=1, value=t)
    ws.cell(row=r, column=2, value=f'=COUNTIFS({col("B")},"{t}",{col("F")},"<>")')
    ws.cell(row=r, column=3, value=f'=COUNTIFS({col("B")},"{t}",{col("AB")},"Sim")')
    c = ws.cell(row=r, column=4, value=f'=IFERROR(C{r}/B{r},"")'); c.number_format = "0.0%"
    ws.cell(row=r, column=5, value=f'=COUNTIFS({col("B")},"{t}",{col("N")},"NC")')
    ws.cell(row=r, column=6, value=f'=COUNTIFS({col("B")},"{t}",{col("R")},"NC")')
    paint(ws, r, r, 2, 6, CALC)

wb.save(OUT)
print(f"OK: {OUT}")
