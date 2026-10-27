# -*- coding: utf-8 -*-
"""
Gera Roteiro_Campo_Fibra.docx: roteiro de campo do estudo de caso
"ocupação de postes da rede de distribuição por cabos de telecomunicações".

Uso:  python3 gerar_roteiro_fibra.py      Requer: pip install python-docx

Valores de referência (conferir no texto oficial antes da redação final):
ABNT NBR 15214:2005 conforme reproduzida em normas de distribuidoras —
5,00 m (ruas e avenidas), 4,50 m (entradas de prédios / uso restrito a
veículos), 6,00 a 7,00 m (rodovias/ferrovias, varia por fonte), 600 mm entre
telecom e neutro, caixa de emenda/reserva técnica no vão a ≥ 2 m do poste,
diâmetro ≤ 65 mm por ponto de fixação, faixa de ocupação de 500 mm.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "Roteiro_Campo_Fibra.docx"
NAVY = RGBColor(0x16, 0x23, 0x3F)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(4)
for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
    h = doc.styles[name]
    h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = NAVY
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(4)
for sec in doc.sections:
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    sec.left_margin = sec.right_margin = Cm(1.8)

fp = doc.sections[0].footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Roteiro de campo — cabos de telecom em postes · pág. "); r.font.size = Pt(8)
run = fp.add_run(); run.font.size = Pt(8)
for el, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                        ("w:instrText", {"xml:space": "preserve"}, "PAGE"),
                        ("w:fldChar", {"w:fldCharType": "end"}, None)):
    n = OxmlElement(el)
    for k, v in attrs.items():
        n.set(qn(k), v)
    if text:
        n.text = text
    run._r.append(n)


def _runs(par, text, size=None, italic=False, bold=False):
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            rr = par.add_run(chunk)
            rr.bold = bold or (i % 2 == 1); rr.italic = italic
            if size:
                rr.font.size = Pt(size)


def p(text, bold=False, size=None, italic=False, align=None, after=None):
    par = doc.add_paragraph()
    _runs(par, text, size, italic, bold)
    if align:
        par.alignment = align
    if after is not None:
        par.paragraph_format.space_after = Pt(after)
    return par


def bullets(items, style="List Bullet"):
    for it in items:
        _runs(doc.add_paragraph(style=style), it)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def fix_widths(t, widths):
    t.autofit = False
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def spacer():
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0)
    sp.add_run().font.size = Pt(4)


def table(header, rows, widths, font=9, blank_rows=0, row_height=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(font)
        shade(c, "DCE3EE")
    for row in list(rows) + [[""] * len(header) for _ in range(blank_rows)]:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            _runs(cells[i].paragraphs[0], str(val), size=font)
    fix_widths(t, widths)
    if row_height:
        for row in t.rows[1:]:
            row.height = Cm(row_height)
    spacer()
    return t


def field_lines(labels, width=17.4):
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for lab in labels:
        cells = t.add_row().cells
        cells[0].text = ""
        rr = cells[0].paragraphs[0].add_run(lab); rr.font.size = Pt(9); rr.bold = True
        shade(cells[0], "F2F4F7")
        t.rows[-1].height = Cm(0.66)
    fix_widths(t, [6.2, width - 6.2])
    spacer()


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


SN = "S / N"

# ============================================================== CAPA =========
p("ROTEIRO DE CAMPO", bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
p("Estudo de caso: ocupação de postes da rede de distribuição por cabos de "
  "telecomunicações", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p("Monografia — Especialização em Engenharia de Segurança do Trabalho "
  "(PECE/EPUSP) — Rodrigo Zambon — tema registrado pela supervisão",
  size=9.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

doc.add_heading("Como usar este roteiro", level=2)
bullets([
    "Leia as seções 1 a 3 antes do campo. **Imprima as fichas** (seção 4): a "
    "**F1** uma por trecho percorrido, a **F2** uma por poste (ou use a "
    "planilha direto no celular), a **F3** sempre que encontrar uma equipe "
    "trabalhando no poste, a **F4** uma por eletricista entrevistado, a **F5** "
    "acompanha todo o estudo.",
    "**Sem nomes nas fichas.** Postes por código (T1-01, T1-02… = trecho 1, "
    "poste 1); trechos por código (T1, T2…); eletricistas por E1, E2…; "
    "provedores por letra (Provedor A, B…). Guarde a relação dos códigos só com você.",
    "Lance os dados em **Planilha_Dados_Fibra.xlsx** (mesma pasta): ela "
    "compara cada medição com o limite da norma e faz os totais.",
    "Antes do primeiro dia, peça a autorização do **Anexo A**.",
])

doc.add_heading("O que o estudo vai mostrar", level=2)
table(["Grupo exposto", "Risco", "O que você vai levantar", "Ficha"],
      [["Eletricistas da distribuidora", "Subida, posicionamento e ancoragem "
        "em poste congestionado; cabos de telecom no caminho; contato "
        "indireto com a rede (NR-10, NR-35)",
        "Nº de cabos e pontos de fixação por poste, equipamentos e sobras no "
        "poste, obstrução da subida; entrevista com eletricistas", "F2, F4"],
       ["Instaladores de telecom", "Choque elétrico perto da rede energizada; "
        "escada apoiada nos cabos no meio da via; falta de sinalização viária",
        "Afastamento telecom–rede elétrica; observação de equipes em serviço",
        "F2, F3"],
       ["População (motociclistas, veículos)", "Cabo baixo, solto ou caído "
        "na via", "Altura do cabo ao solo no meio do vão; cabos soltos, "
        "caídos ou abandonados", "F2"]],
      [3.6, 5.4, 6.2, 2.2])

doc.add_heading("Regras de segurança para você", level=2)
bullets([
    "**Tudo é medido do chão.** Não suba em poste, escada ou veículo para medir.",
    "**Nunca use trena metálica, vara ou haste perto da rede elétrica**, nem "
    "encoste em cabo de telecom: ele pode estar energizado por contato com a rede.",
    "**Cabo caído no chão: não toque.** Isole com distância, fotografe de "
    "longe e avise a distribuidora (risco grave e iminente).",
    "Fique na calçada; use **colete refletivo**; atenção ao trânsito ao "
    "medir no meio da rua; prefira horários de pouco movimento e luz do dia.",
    "Ao observar equipes de telecom, **não interfira** e não se aproxime da "
    "área de trabalho. Fotografe sem rostos e sem logotipos.",
])

page_break()

# ============================================================ 1 PREPARAÇÃO ==
doc.add_heading("1  Antes do campo — preparação", level=1)
table(["#", "Tarefa", "Quando", "Feito"],
      [["1", "Autorização por escrito da distribuidora para o levantamento e "
             "uso de informações (Anexo A)", "antes da 1ª saída", "☐"],
       ["2", "Pedir à distribuidora (anonimizado): norma técnica própria de "
             "compartilhamento de postes, nº de empresas ocupantes, "
             "notificações de irregularidade, registros de ocorrências "
             "envolvendo cabos de telecom", "antes da 1ª saída", "☐"],
       ["3", "Conseguir o instrumento de medição de altura (ver 1.1)", "antes da 1ª saída", "☐"],
       ["4", "Escolher os trechos (ver 1.2) e definir os códigos T1, T2…", "antes da 1ª saída", "☐"],
       ["5", "Imprimir fichas; colete refletivo, prancheta, trena comum "
             "(para distâncias no chão), celular carregado com GPS", "véspera", "☐"],
       ["6", "Saída piloto: medir 5 postes para testar o método e as fichas", "1ª saída", "☐"],
       ["7", "Agendar 4 a 6 entrevistas com eletricistas (F4)", "novembro", "☐"]],
      [0.8, 11.6, 3.4, 1.4])

doc.add_heading("1.1  Como medir a altura dos cabos do chão", level=2)
table(["Opção", "Como funciona", "Observação"],
      [["**Medidor ultrassônico de altura de cabos** (recomendado)",
        "Aparelho de mão apontado para cima, embaixo do vão; mede a altura de "
        "vários cabos de uma vez", "Feito para isso; alugável ou emprestado "
        "por empresas de telecom/eletricidade. Anotar modelo e precisão."],
       ["**Trena a laser + método do ângulo**",
        "Mede a distância horizontal até a base do poste ou um ponto sob o "
        "cabo e o ângulo de visada (app de inclinômetro); altura = distância "
        "× tangente(ângulo) + altura dos olhos", "Mais barato; erro maior "
        "(registrar). A planilha calcula se você lançar distância e ângulo."],
       ["**Referência na foto**", "Foto com referência de altura conhecida "
        "(ex.: placa de rua padrão, porta)", "Só para casos extremos "
        "(cabo visivelmente baixo); anotar “estimado”."]],
      [4.2, 7.2, 6.0])
p("**Onde medir:** a altura que conta é a do **ponto mais baixo do cabo**, "
  "normalmente o **meio do vão** entre dois postes, não junto ao poste. Anote "
  "também a temperatura do ar: a norma considera a flecha máxima (cabo "
  "aquecido); medir em dia fresco **favorece** o cabo, então uma não "
  "conformidade encontrada é conservadora.", size=9.5)

doc.add_heading("1.2  Amostra", level=2)
bullets([
    "Meta: **60 a 100 postes consecutivos** (com o vão seguinte), em 4 trechos:",
    "T1 — centro/avenida comercial; T2 — bairro residencial; T3 — via com "
    "travessia de rua movimentada; T4 — trecho rural ou de estrada.",
    "Em cada trecho, postes **seguidos**, sem escolher “os piores”: isso dá "
    "valor estatístico (percentual de postes irregulares).",
])

doc.add_heading("1.3  Valores de referência (conferir antes da redação)", level=2)
table(["Requisito", "Valor", "Fonte"],
      [["Altura mínima do cabo ao solo — ruas e avenidas", "5,00 m", "NBR 15214"],
       ["Altura mínima — entradas de prédios / locais de uso restrito a veículos", "4,50 m", "NBR 15214"],
       ["Altura mínima — travessia de rodovia/ferrovia", "6,00 m a 7,00 m (varia por fonte)", "NBR 15214 / normas de distribuidoras"],
       ["Afastamento mínimo entre telecom e o neutro da rede elétrica", "0,60 m", "NBR 15214"],
       ["Caixa de emenda e reserva técnica (sobra)", "no vão, ≥ 2,0 m do poste", "NBR 15214"],
       ["Diâmetro máximo do conjunto por ponto de fixação", "65 mm", "NBR 15214"],
       ["Faixa de ocupação reservada à telecom no poste", "500 mm", "NBR 15214"],
       ["Identificação dos cabos e equipamentos pela ocupante", "obrigatória", "NBR 15214 / Res. Conj. ANEEL-Anatel nº 4/2014"]],
      [9.6, 4.0, 3.8])
p("Se a distribuidora tiver **norma própria de compartilhamento**, ela "
  "prevalece no estudo de caso (ajuste os limites na aba Referencia da planilha).",
  italic=True, size=9)

doc.add_heading("1.4  Agenda", level=2)
table(["Período", "O quê"],
      [["novembro", "Autorização, instrumento, saída piloto, trechos T1 e T2"],
       ["até meados de dezembro", "Trechos T3 e T4; entrevistas (F4); observação de equipes (F3)"],
       ["até 20/12", "Enviar planilha/fichas, fotos e documentos da distribuidora"],
       ["dez. a jan.", "Redação dos capítulos de métodos, resultados e discussão"],
       ["até 22/02/2027", "Entrega da monografia em PDF no Moodle"]],
      [4.5, 12.9])

page_break()

# ============================================================ 2 PROTOCOLO ===
doc.add_heading("2  Protocolo por poste (passo a passo)", level=1)
for s in [
    "Dê o código ao poste (T1-01…) e marque o ponto de GPS no celular. "
    "(Coordenadas ficam só nos seus dados; no TCC aparecem apenas trechos.)",
    "Classifique o **local do vão seguinte**: rua/avenida, entrada de "
    "veículos, travessia, rodovia, só pedestres.",
    "Identifique a **estrutura elétrica**: primário + secundário, só "
    "primário, só secundário, e se há transformador ou chave no poste.",
    "**Conte** os cabos de telecom e os **pontos de fixação** (pontos de "
    "amarração no poste).",
    "**Meça** a altura do cabo de telecom **mais baixo no meio do vão** e "
    "a altura do condutor elétrico mais baixo (neutro) e do cabo de telecom "
    "mais alto **junto ao poste**; a planilha calcula o afastamento.",
    "Verifique: plaqueta de identificação; sobra de cabo (reserva técnica) "
    "enrolada no poste; caixa de emenda a menos de 2 m do poste; "
    "equipamento ativo (fonte, amplificador) no poste; cabo solto, "
    "pendurado, caído ou abandonado (sem uso aparente); cabo encostando em "
    "condutor elétrico ou em ferragem; poste danificado ou inclinado.",
    "Avalie a **obstrução da subida** (baixa / média / alta): espaço livre "
    "para a escada e para o eletricista passar entre os cabos até a rede.",
    "Fotografe: poste inteiro, vão, detalhe das irregularidades (F5).",
]:
    _runs(doc.add_paragraph(style="List Number"), s)
p("Tempo estimado: 5 a 8 minutos por poste — cerca de 15 a 20 postes por "
  "manhã.", italic=True, size=9.5)

doc.add_heading("2.1  Como os resultados serão apresentados (a planilha calcula)", level=2)
bullets([
    "% de postes com **pelo menos uma** não conformidade, por trecho e no total;",
    "% de vãos com cabo **abaixo da altura mínima**, e a menor altura encontrada;",
    "% de postes com **afastamento telecom–neutro abaixo de 0,60 m**;",
    "média e máximo de cabos e de pontos de fixação por poste;",
    "frequência de cada irregularidade (sobra no poste, caixa de emenda, "
    "cabo caído, sem identificação…);",
    "relação entre a quantidade de cabos e a obstrução da subida (risco ao eletricista).",
])

page_break()

# ============================================================== FICHAS =====
doc.add_heading("Fichas de campo", level=1)
doc.add_heading("F1 — Identificação do trecho", level=2)
field_lines(["Trecho (código T1, T2…)", "Data / horário de início e fim",
             "Tipo (centro, residencial, travessia, rural/estrada)",
             "Temperatura do ar e tempo (sol, nublado, vento)",
             "Instrumento de altura: tipo, modelo, precisão",
             "Postes levantados (do código … ao código …)",
             "Observações gerais do trecho"])

doc.add_heading("F2 — Poste (uma por poste)", level=2)
field_lines(["Código do poste / GPS anotado?",
             "Local do vão seguinte (rua-avenida / entrada de veículos / "
             "travessia / rodovia / só pedestres)",
             "Estrutura elétrica (prim.+sec. / só prim. / só sec.; trafo ou chave?)",
             "Nº de cabos de telecom / nº de pontos de fixação",
             "Altura do cabo de telecom mais baixo no MEIO DO VÃO (m)",
             "Altura do neutro (condutor elétrico mais baixo) no poste (m)",
             "Altura do cabo de telecom mais alto no poste (m)",
             "Plaqueta de identificação em todos os cabos? (S/N/parcial)",
             "Sobra de cabo enrolada no poste? (S/N)",
             "Caixa de emenda a menos de 2 m do poste? (S/N)",
             "Equipamento ativo no poste (fonte, amplificador)? (S/N)",
             "Cabo solto, pendurado, caído ou abandonado? (descrever)",
             "Cabo encostando em condutor elétrico/ferragem? (S/N)",
             "Poste danificado ou inclinado? (S/N)",
             "Obstrução da subida (baixa / média / alta)",
             "Nº das fotos"])

page_break()
doc.add_heading("F3 — Observação de equipe de telecom em serviço", level=2)
p("Preencha só quando encontrar uma equipe trabalhando. Observe de longe; não "
  "interfira; fotos sem rosto e sem logotipo.", size=9.5)
field_lines(["Data / horário / trecho (código)", "Provedor (letra A, B…) / nº de pessoas",
             "Atividade (instalação, manutenção, retirada de cabo)"])
table(["#", "Item observado", "Resultado", "Obs."],
      [["1", "Sinalização viária (cones, placas, isolamento) quando a escada/veículo ocupa a via", "C / NC / NA", ""],
       ["2", "Escada apoiada no POSTE (não nos cabos) e amarrada/estabilizada", "C / NC / NA", ""],
       ["3", "Escada de material isolante (fibra), sem partes metálicas", "C / NC / NA", ""],
       ["4", "Trabalhador conectado (cinturão + talabarte) acima de 2 m", "C / NC / NA", ""],
       ["5", "Pessoa no solo apoiando a escada / controlando o trânsito", "C / NC / NA", ""],
       ["6", "Trabalhador mantendo distância da rede elétrica (fora da faixa da rede)", "C / NC / NA", ""],
       ["7", "Capacete com jugular, luvas, calçado de segurança", "C / NC / NA", ""],
       ["8", "Trabalho suspenso em chuva/vento", "C / NC / NA", ""],
       ["9", "Cabo novo instalado dentro da faixa de ocupação e com plaqueta", "C / NC / NA", ""]],
      [0.8, 11.2, 2.4, 3.0], font=8.5)

doc.add_heading("F4 — Conversa com eletricista da distribuidora (anônima, voluntária)", level=2)
p("Código: ______   Data: ____/____/______   Tempo na função: ______", size=9.5)
qs = ["Os cabos de telecom atrapalham a subida e o posicionamento no poste? Como?",
      "Onde você apoia a escada / ancora o talabarte quando o poste está cheio de cabos?",
      "Já passou por situação de risco por causa desses cabos (choque, queda, "
      "quase acidente, cabo enroscando)? Conte sem nomes.",
      "Já encontrou cabo de telecom encostando na rede ou energizado?",
      "Já viu instalador de telecom em situação perigosa (escada nos cabos, sem "
      "sinalização, perto da rede)?",
      "Já atendeu ocorrência com cabo caído/baixo atingindo pessoa ou veículo?",
      "Quanto tempo a mais leva um serviço em poste congestionado?",
      "O que a empresa poderia fazer para reduzir esse risco?"]
for i, q in enumerate(qs, 1):
    p(f"{i}. {q}", size=9.5, after=1)
    for _ in range(2):
        p("_" * 98, size=8, after=0)

page_break()
doc.add_heading("F5 — Registro fotográfico", level=2)
p("Fotos são **obrigatórias no pôster**. Regras: sem rostos, sem logotipos de "
  "empresas, **sem nomes de provedores legíveis nas plaquetas**, sem placas de "
  "veículos e sem fachadas/nomes de rua identificáveis. Se não der para evitar, "
  "fotografe e marque “desfocar” — o desfoque é feito antes de entrar no TCC.",
  size=9.5)
p("Fotos que precisamos ter:", bold=True, size=9.5)
bullets(["poste congestionado (visão geral)", "vão com cabo baixo sobre a via",
         "sobra de cabo enrolada no poste", "caixa de emenda junto ao poste",
         "cabo caído ou solto", "cabo de telecom muito próximo da rede elétrica",
         "equipe de telecom: escada nos cabos / via sem sinalização",
         "medição sendo feita (instrumento em uso)"])
table(["Nº foto", "Data", "Poste/trecho", "O que mostra", "Desfocar?"],
      [], [1.6, 2.0, 2.8, 8.2, 2.8], blank_rows=14, row_height=0.65)

page_break()
doc.add_heading("3  Depois do campo — o que me enviar", level=1)
bullets(["planilha preenchida (ou fotos legíveis das fichas);",
         "fotos com a lista de numeração;",
         "respostas das entrevistas (F4) e observações de equipes (F3);",
         "documentos da distribuidora: norma de compartilhamento, nº de "
         "ocupantes, notificações e registros de ocorrências (anonimizados);",
         "modelo e precisão do instrumento de medição."])
p("Com isso eu escrevo os capítulos de métodos, resultados e discussão, monto "
  "tabelas e gráficos e atualizo pôster e apresentação.")

doc.add_heading("Anexo A — Modelo de autorização da distribuidora", level=1)
p("AUTORIZAÇÃO PARA REALIZAÇÃO DE ESTUDO DE CASO", bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER)
p("A [nome da organização], inscrita no CNPJ [nº], autoriza [nome completo "
  "do aluno], aluno do curso de Especialização em Engenharia de Segurança do "
  "Trabalho do Programa de Educação Continuada em Engenharia da Escola "
  "Politécnica da Universidade de São Paulo, a realizar estudo de caso para "
  "sua monografia, compreendendo: (a) levantamento visual e medições, a partir "
  "do solo, da ocupação de postes de sua rede de distribuição por cabos de "
  "telecomunicações; (b) entrevistas voluntárias e anônimas com eletricistas; "
  "(c) consulta a normas técnicas, registros de ocupantes, notificações e "
  "ocorrências relacionadas ao compartilhamento de postes; e (d) registro "
  "fotográfico.")
p("A autorização fica condicionada a que o trabalho não identifique a "
  "organização, seus trabalhadores, as empresas ocupantes ou terceiros, "
  "inclusive em fotografias, e que as informações sejam usadas exclusivamente "
  "para fins acadêmicos.")
p("[Cidade], ____ de ______________ de 2026.")
p("\n\n______________________________________\n[Nome e cargo do responsável]")

doc.save(OUT)
print(f"OK: {OUT}")
